# Import libraries
from docxtpl import DocxTemplate
from docx.shared import Inches
from docx import Document
from docx2pdf import convert

import pyodbc
import pandas as pd
import os
import glob
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
import matplotlib.ticker as ticker
import datetime
import re
import sys
from pathlib import Path
import shutil
import tempfile
import babel.numbers


def read_csv_file(data_filepath):
    df = pd.read_csv(data_filepath, sep="|")
    print("***** CSV file read successfully *****\n")
    return df


def row_picker(data_filepath):
    # Step 1: Import csv file
    pd.set_option('display.max_columns', None) #show all displayed columns
    df = read_csv_file(data_filepath).fillna("-", inplace=False)
    df_display = pd.DataFrame(df.get(['rn', 'refdate', 'trgRarRank', 'trgOCF', 'ISIN', 'Platform', 'FundName']))

    # Step 2: Display rows and ask for user input
    print('\n===================================================================================================================')
    print('Please input rn to be processed and press Enter. No space in between. E.g. 1,2-8,10.')
    print('To publish only classes with breaches, input 0 and press Enter.')
    print('===================================================================================================================')
    print(df_display.to_string(index=False))

    userinput = input('\nInput here:')
    print(f'You have inputted: {userinput}')

    # Step 3: Decompose input
    if userinput == '0': #if publishing breached classes only
        filtered_df = df.loc[(df['trgRarRank'] == 1) | (df['trgOCF'] == 1), ['rn']]
        final_list = filtered_df['rn'].to_list()

        if not final_list: #blank list
            print("***** There are no breaches. No publication required. *****\n")
            sys.exit(0)
        else:
            print("***** These rn would be executed: " + str(final_list) + " *****\n")
            return final_list, df

    else:
        inter_list = userinput.split(',')
        final_list = []

        for i in inter_list:
            if "-" in i:
                start_pattern = r"\d+-"
                end_pattern = r"-\d+"
                start_number = int(re.search(start_pattern, i).group().replace("-",""))
                end_number = int(re.search(end_pattern, i).group().replace("-",""))
                for r in range(start_number,end_number+1):
                    final_list.append(r)
                #Note that range doesn't include end number.
            else:
                final_list.append(int(i))

        final_list = list(dict.fromkeys(final_list)) #Unique values
        # print(final_list)

        # Step 4: Verify input
        df_rn = pd.Series(df.get('rn')).tolist()
        # print(df_rn)
        # print([i for i in final_list if i not in df_rn])
        # print(len([i for i in final_list if i not in df_rn]))
        if len([i for i in final_list if i not in df_rn]) != 0: #Check if there are any out scope numbers. Reject if any.
            raise ValueError("***** Terminated due to out of scope rn: " + str([i for i in final_list if i not in df_rn]) + " *****\n")
        else:
            print("***** These rn would be executed: " + str(final_list) + " *****\n")
            return final_list, df


def word_to_pdf(outputpath, wordoutputpath, pdf_filename):
    print('Word file exists: ', os.path.exists(wordoutputpath))
    # Copy word file to temp directory and run on the copy to ensure PDF function can be run
    # This also shortens path to avoid the word file name length limit
    with tempfile.TemporaryDirectory() as tmp_dir:
        local_docx = os.path.join(tmp_dir, os.path.basename(wordoutputpath))
        local_pdf = os.path.join(tmp_dir, f"{pdf_filename}.pdf")
        shutil.copy(wordoutputpath, local_docx)

        convert(local_docx, local_pdf)

        final_pdf = os.path.join(outputpath, f"{pdf_filename}.pdf")
        shutil.copy(local_pdf, final_pdf)

    print("***** Word file saved as PDF file successfully *****\n")


def fname_indexer(folderpath, core, ftype):
    ncounter = 0
    fname = core + f"{ncounter:03}" + "." + ftype

    while glob.glob(os.path.join(folderpath, fname), recursive=True):
        # print(ncounter, " ncounter doesn't work. Next.")
        ncounter += 1
        fname = core + f"{ncounter:03}" + "." + ftype

    # print(ncounter, " ncounter works.")
    fpath = os.path.join(folderpath, fname)

    print(fpath)
    return fpath


def generate_cci_summary(df, rn, wdtemplate_path, pdffolder_path):
    # Step 1: Convert dataframe to list of dictionaries
    # df = read_csv_file(datasheet_path).fillna("-", inplace=False)
    data_dict = df.to_dict(orient='records')

    # Step 2: Define file link for CCI Summary word template
    doc = DocxTemplate(wdtemplate_path)
    print("***** File Link for CCI Summary Word Template defined successfully *****\n")

    # Step 3: Match the bits
    ### 3a. Determine the row in data_dict to process
    r = int(rn) - 1 # Need to minus 1 from rn inputted by user to convert it to corresponding row in data_dict.
    datarow = data_dict[r]

    print("***** Data row defined as " + str(r) + " successfully *****\n")

    ### 3b. For entity descriptions
    if datarow['FundType'] == 'RIAIF':
        pro_type_stylised = 'a Retail Investor Alternative Investment Fund (RIAIF), ' \
                            'an open-ended Irish collective asset management vehicle which is constituted as an umbrella fund with segregated liability between sub-funds, with variable capital'
    elif datarow['FundType'] == 'UCITS':
        pro_type_stylised = 'a Undertakings for the Collective Investment in Transferable Securities (UCITS), ' \
                            'an open-ended Irish collective asset management vehicle which is constituted as an umbrella fund with segregated liability between sub-funds, with variable capital'
    print("***** Entity descriptions defined successfully *****\n")

    ### 3c. For rep class table
    ### N.B. to_dict converts the whole dataframe to a LIST of DICTIONARIES and we need to pick one of them for one row
    isin_repclassdf = df.filter(regex=r'ISIN_RepBy\w+')
    isin_repclassdict = isin_repclassdf.to_dict(orient='records')

    class_repclassdf = df.filter(regex=r'Class_RepBy\w+')
    class_repclassdict = class_repclassdf.to_dict(orient='records')

    entryFee_repclassdf = df.filter(regex=r'entryFee_RepBy\w+')
    entryFee_repclassdict = entryFee_repclassdf.to_dict(orient='records')
    print("***** Rep class details extracted from dataframe successfully *****\n")

    repclassdict = {}
    repclasslist = []
    for i in range(len(class_repclassdict[r])):
        if list(class_repclassdict[r].values())[0] == "-":  # Where there are no classes being represented
            repclassdict = {'class_name': 'Not applicable',
                            'isin': 'Not applicable',
                            'entry_fee': 'Not applicable'}
            repclasslist.append(repclassdict)
            break
        elif list(class_repclassdict[r].values())[0] != "-":
            try:
                repclassdict = {'class_name': list(class_repclassdict[r].values())[i],
                                'isin': list(isin_repclassdict[r].values())[i],
                                'entry_fee': f"{float(list(entryFee_repclassdict[r].values())[i]) * 100:.2f}%"}
            except ValueError: # when "-" is met as all non-null values are processed
                break
            repclasslist.append(repclassdict)
    # print(repclasslist)
    print("***** Rep class dictionary list built successfully *****\n")

    ### 3d. For rank graphics
    ranklist = []
    rankdict = {}
    for i in range(10):
        if i + 1 == datarow['AdjScore']:
            rankdict = {'rank': i + 1, 'bg': 'BF9A5E'}
        else:
            rankdict = {'rank': i + 1, 'bg': '380C57'}
        ranklist.append(rankdict)
    # print(ranklist)
    print("***** Rank dictionary list built successfully *****\n")

    ### 3e. For performance descriptions
    ### N.B. to_dict converts the whole dataframe to a LIST of DICTIONARIES and we need to pick one of them for one row
    date_chgdf = df.filter(regex=r'refdate_chg\w+')
    date_chgdict = date_chgdf.to_dict(orient='records')

    desc_chgdf = df.filter(regex=r'^chg\w+')
    desc_chgdict = desc_chgdf.to_dict(orient='records')

    # print(date_chgdict)
    # print(desc_chgdict)
    print("***** Performance descriptions extracted from dataframe successfully *****\n")

    chgdict = {}
    chglist = []
    for i in range(len(date_chgdict[r])):
        if list(date_chgdict[r].values())[0] == "-":  # Where there are no descriptions
            chgdict = {'date': '',
                       'desc': ''}
            break
        elif list(date_chgdict[r].values())[0] != "-":
            try:
                if list(desc_chgdict[r].values())[i] == "-": # Ensure no "-" appears in the document
                    chgdict = {'date': '',
                               'desc': ''}
                else:
                    chgdict = {'date': list(date_chgdict[r].values())[i],
                               'desc': list(desc_chgdict[r].values())[i]}
            except ValueError: # in case of errors
                break
            chglist.append(chgdict)
    # print(chglist)
    print("***** Performance description dictionary list built successfully *****\n")

    ### 3f. For end date of performance period
    yearfund_df = df.filter(regex=r'^y\d+$')
    yearfund_dict = yearfund_df.to_dict(orient='records')

    end_date = max(
        [datetime.datetime.strptime(x, "%Y-%m-%d") for x in list(yearfund_dict[r].values()) if x != "-"]).date()
    # print(end_date)
    print("***** end_date extracted successfully *****\n")

    # Step 4: Commit context
    context = {
        'Product_Name_Fund': datarow['Portfolio_Name'],
        'Product_Umbrella': datarow['Umbrella'],
        'Product_Name_Class': datarow['FundName'],
        'Product_ID': datarow['ISIN'],
        'Product_Type': pro_type_stylised,
        'IM_Name': datarow['InvF_Name'],
        'FCA_Authorised': datarow['InvF_FCA_Auth'],
        'Manufacturer_Name': datarow['Mfr_Name'],
        'Mfr_Desc1': datarow['Mfr_Desc1'],
        'Fund_Mgt_Type': datarow['MgtType'],
        'Fund_Mgt_Name': datarow['MgtF_Name'],
        'MgtF_Desc1': datarow['MgtF_Desc1'],
        'Ref_Date': datetime.datetime.strftime(datetime.datetime.strptime(datarow['refdate'], "%Y-%m-%d").date(),"%d/%m/%Y"),
        'Inv_Obj_1': datarow['InvObj1'],
        'Inv_Obj_2': datarow['InvObj2'],
        'Benchmark_Name': datarow['BM_Name_Stylised'],
        'repclasslist': repclasslist,
        'RHP': datarow['RHP'],
        'Base_Currency': datarow['Base_Currency'],
        'Manufactured_Year': datetime.datetime.strftime(datetime.datetime.strptime(datarow['estYear'], "%Y-%m-%d").date(),"%d/%m/%Y"),
        'ranklist': ranklist,
        'start_date': datetime.datetime.strftime(datetime.datetime.strptime(datarow['y0'], "%Y-%m-%d").date(),"%d/%m/%Y"),
        'end_date': datetime.datetime.strftime(end_date,"%d/%m/%Y"),
        'chglist': chglist,
        'Assumed_Inv_Amt': babel.numbers.format_currency(datarow['assumed_inv_amt'],datarow['Base_Currency'],u'¤#,##0',locale='en_US',currency_digits=False),
        'Avg_Return': str(round(datarow['median_ret'] * 100, 2)) + "%",
        'One_Off_Entry_Cost_Percent': f"{float(datarow['entry_p']) * 100:.2f}%",
        'One_Off_Exit_Cost_Percent': f"{float(datarow['exit_p']) * 100:.2f}%",
        'One_Off_Entry_Cost_Abs': babel.numbers.format_currency(datarow['entry_abs'],datarow['Base_Currency'],u'¤#,##0',locale='en_US',currency_digits=False),
        'One_Off_Exit_Cost_Abs': babel.numbers.format_currency(datarow['exit_abs'],datarow['Base_Currency'],u'¤#,##0',locale='en_US',currency_digits=False),
        'Ongoing_Cost_Percent': f"{float(datarow['oc_p']) * 100:.2f}%",
        'Closed_Ended_Ongoing_Cost_Percent': f"{float(datarow['ce_oc_p']) * 100:.2f}%",
        'Trans_Cost_Percent': f"{float(datarow['explicit_tc']) * 100:.2f}%",
        'Ongoing_Cost_Abs': babel.numbers.format_currency(datarow['oc_abs'],datarow['Base_Currency'],u'¤#,##0',locale='en_US',currency_digits=False),
        'Closed_Ended_Ongoing_Cost_Abs': babel.numbers.format_currency(datarow['ce_oc_abs'],datarow['Base_Currency'],u'¤#,##0',locale='en_US',currency_digits=False),
        'MgtF_Address': datarow['MgtF_Address'],
        'MgtF_Email': datarow['MgtF_Email']
    }

    # Step 5: Render document
    doc.render(context)

    # Step 6: Save the generated file
    wordoutputname = datetime.datetime.strftime(datetime.datetime.strptime(datarow['refdate'],"%Y-%m-%d"),"%Y%m") + \
                 " " + \
                 datarow['FundName']

    wordoutputpath = os.path.join(pdffolder_path, wordoutputname + ".docx")

    with tempfile.TemporaryDirectory() as tmp_dir: #Save locally and move file back to circumvent word file name length limit
        local_docx = Path(tmp_dir) / Path(wordoutputpath).name
        doc.save(local_docx)
        shutil.copy(local_docx, wordoutputpath)

    print("***** Word file: " + wordoutputname + " is generated *****\n")

    return wordoutputname, wordoutputpath


def insert_line_graph(df, row, plotfolder_path, wordoutputpath):
    # Step 1: Create and save the plot
    ### 1a. Create lists for line graph
    ### N.B. to_dict converts the whole dataframe to a LIST of DICTIONARIES and we need to pick one of them for one row
    r = row - 1 # To convert rn to row number
    years_df = df.filter(regex=r'^y\d+$')
    years_dict = years_df.to_dict(orient='records')
    xlist = [datetime.datetime.strptime(x,"%Y-%m-%d") for x in list(years_dict[r].values()) if x != '-']

    fundamt_df = df.filter(regex=r'^y\d+_amt$')
    fundamt_dict = fundamt_df.to_dict(orient='records')
    y_fund_list = [round(float(x),0) for x in list(fundamt_dict[r].values()) if x != '-']

    bmamt_df = df.filter(regex=r'^y\d+_amt_bm$')
    bmamt_dict = bmamt_df.to_dict(orient='records')
    y_bm_list = [round(float(x),0) for x in list(bmamt_dict[r].values()) if x != '-']

    ### 1b. Plot the graph
    ### Create figure and axes with transparent background
    fig = plt.figure()
    ax = fig.add_axes((0.1, 0.1, 0.8, 0.8))
    ax.plot(xlist, y_fund_list, label='Fund', color='gold')
    ax.plot(xlist, y_bm_list, label='Benchmark', color='grey', linestyle='--')

    # plt.xlabel('Year')
    plt.ylabel('Value in ' + df.get('Base_Currency')[r])

    date_format = mdates.DateFormatter('%d/%m/%Y')
    ax.xaxis.set_major_formatter(date_format)
    ax.xaxis.set_major_locator(mdates.AutoDateLocator(maxticks=11))
    fig.autofmt_xdate()

    ax.yaxis.set_major_formatter(ticker.StrMethodFormatter('{x:,.0f}'))


    plt.legend(loc='upper left')

    ### Add horizontal grid lines
    plt.grid(axis='y')

    ### Remove axes background (transparent)
    ax.set_facecolor('none')

    ### Remove figure background (transparent)
    fig.patch.set_alpha(0.0)

    ### Determine the plot image file name
    fname = datetime.datetime.strftime(datetime.datetime.strptime(df.get('refdate')[r],"%Y-%m-%d"),"%Y%m") + \
                 df.get('ISIN')[r] + "_"
    fpath = fname_indexer(plotfolder_path, fname, "png")
    plt.savefig(fpath, transparent = True, dpi=300, bbox_inches='tight')
    plt.close(fig)

    print("***** Line graph plotted and saved successfully *****\n")

    # Step 2: Open the Word document
    docp = Document(wordoutputpath)

    # Step 3: Find placeholder text and insert image after it
    placeholder = "[There is no past performance available.]"
    for table in docp.tables:
        for row in table.rows:
            for cell in row.cells:
                if placeholder in cell.text:
                    # Remove placeholder text
                    cell.text = cell.text.replace(placeholder, "")
                    # Insert image after this paragraph
                    paragraph = cell.paragraphs[0]
                    run = paragraph.add_run()
                    run.add_picture(fpath, width=Inches(4))
                    break

    # Step 4: Save the updated document
    docp.save(str(wordoutputpath))
    print("***** Plot inserted successfully *****\n")


def update_log(df, row, datasheetname):
    # Connection parameters
    server = 'devuksvappimo01\IMDR'
    database = 'Sandbox'

    # Connection string for Windows Authentication
    connection_string = f'DRIVER={{ODBC Driver 17 for SQL Server}};SERVER={server};DATABASE={database};Trusted_Connection=yes'
    connection = pyodbc.connect(connection_string)

    print("***** Connected to SQL database " + server + "\\" + database + " using Windows Authentication *****\n")

    # Execute stored procedure
    cursor = connection.cursor()
    proc_sql = "EXEC ukcci_log_publish_refresh ?, ?, ?, ?"

    r = row - 1 # To convert rn to row number

    refdate = datetime.datetime.strftime(datetime.datetime.strptime(df.get('refdate')[r],"%Y-%m-%d"),"%Y-%m-%d")
    isin = df.get('ISIN')[r]
    alpha_code = df.get('Alpha_Code')[r]

    try:
        cursor.execute(proc_sql, refdate, isin, alpha_code, datasheetname)
        connection.commit()
        print("***** Python sent instruction to execute ukcci_log_publish_refresh *****")
    except pyodbc.Error as e: #Python error
        print(f"Error: {e}")

    print(cursor.messages) #Print out SQL messages
    print("\n")

    # Close connection
    cursor.close()
    connection.close()